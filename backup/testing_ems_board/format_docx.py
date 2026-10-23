import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    # remove existing shd if any
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

def parse_original_document(file_path):
    src_doc = docx.Document(file_path)
    paras = [p.text for p in src_doc.paragraphs]

    # Parse Quiz 1
    q1_list = []
    curr = None
    for p in paras[5:107]:
        text = p.strip()
        if not text:
            continue
        m = re.match(r'^(\d+)\.\s*(.*)', text)
        if m and int(m.group(1)) <= 20:
            if curr:
                q1_list.append(curr)
            curr = {'num': int(m.group(1)), 'q': m.group(2).strip(), 'opts': []}
        elif curr and re.match(r'^[a-d]\)\s*(.*)', text, re.IGNORECASE):
            m_opt = re.match(r'^([a-d])\)\s*(.*)', text, re.IGNORECASE)
            curr['opts'].append((m_opt.group(1).upper(), m_opt.group(2).strip()))
    if curr:
        q1_list.append(curr)

    # Parse Quiz 1 Answer Key
    ans1_list = []
    for p in paras[107:128]:
        text = p.strip()
        if not text or text == 'Answer Key':
            continue
        m = re.match(r'^(\d+)\.\s*(.*)', text)
        if m:
            ans_raw = m.group(2).strip()
            if len(ans_raw) > 1 and ans_raw[0] in 'abcd' and ans_raw[1] in ' )':
                ans_raw = ans_raw[0].upper() + ans_raw[1:]
            ans1_list.append((int(m.group(1)), ans_raw))

    # Parse Quiz 2
    q2_list = []
    curr = None
    for p in paras[129:]:
        text = p.strip()
        if not text:
            continue
        cleaned = re.sub(r'^[\t\s\u2022\-\*]+\s*', '', text)
        m = re.match(r'^(\d+)\.\s*(.*)', cleaned)
        if m and int(m.group(1)) <= 20:
            if curr:
                q2_list.append(curr)
            curr = {'num': int(m.group(1)), 'q': m.group(2).strip(), 'opts': [], 'ans': None}
        elif curr and 'Answer:' in cleaned:
            curr['ans'] = cleaned.split('Answer:', 1)[1].strip()
        elif curr and re.match(r'^[A-D]\)\s*(.*)', cleaned):
            m_opt = re.match(r'^([A-D])\)\s*(.*)', cleaned)
            curr['opts'].append((m_opt.group(1), m_opt.group(2).strip()))
    if curr:
        q2_list.append(curr)

    return q1_list, ans1_list, q2_list

def create_aligned_document(q1_list, ans1_list, q2_list, output_path):
    doc = docx.Document()

    # Page Margins: 1 inch (72 pt) all around
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Palette
    C_NAVY = RGBColor(27, 54, 93)      # #1B365D - Headers / Accent
    C_SLATE = RGBColor(71, 85, 105)    # #475569 - Secondary / Subtitles
    C_DARK = RGBColor(30, 41, 59)      # #1E293B - Questions / Body Text
    C_BODY = RGBColor(51, 65, 85)      # #334155 - Option text
    C_GREEN = RGBColor(22, 101, 52)    # #166534 - Answer text
    HEX_NAVY = "1B365D"
    HEX_ROW_EVEN = "F8FAFC"
    HEX_ROW_ODD = "FFFFFF"

    # Base Normal Style setup
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = C_BODY

    # Main Document Header
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("History & Freedom Fighters Quiz")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = C_NAVY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(20)
    r_sub = p_sub.add_run("A 40-Question Multiple-Choice Assessment with Comprehensive Answer Keys")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = C_SLATE

    # ==========================================
    # PART 1: The History Rebellions Quiz
    # ==========================================
    p_p1 = doc.add_paragraph()
    p_p1.paragraph_format.space_before = Pt(10)
    p_p1.paragraph_format.space_after = Pt(4)
    p_p1.paragraph_format.keep_with_next = True
    r_p1 = p_p1.add_run("Part 1: The History Rebellions Quiz")
    r_p1.font.name = 'Calibri'
    r_p1.font.size = Pt(15)
    r_p1.font.bold = True
    r_p1.font.color.rgb = C_NAVY

    p_desc1 = doc.add_paragraph()
    p_desc1.paragraph_format.space_before = Pt(0)
    p_desc1.paragraph_format.space_after = Pt(14)
    p_desc1.paragraph_format.keep_with_next = True
    r_desc1 = p_desc1.add_run(
        "Test your knowledge of major historical rebellions with this 20-question multiple-choice quiz covering "
        "famous uprisings, leaders, and turning points from world and Indian history. Answers are provided in the key at the end of this section."
    )
    r_desc1.font.name = 'Calibri'
    r_desc1.font.size = Pt(10)
    r_desc1.font.italic = True
    r_desc1.font.color.rgb = C_SLATE

    for item in q1_list:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.left_indent = Inches(0)
        p_q.paragraph_format.space_before = Pt(9)
        p_q.paragraph_format.space_after = Pt(3)
        p_q.paragraph_format.line_spacing = 1.15
        p_q.paragraph_format.keep_with_next = True
        
        r_qnum = p_q.add_run(f"{item['num']}. ")
        r_qnum.font.name = 'Calibri'
        r_qnum.font.size = Pt(11)
        r_qnum.font.bold = True
        r_qnum.font.color.rgb = C_NAVY

        r_qtxt = p_q.add_run(item['q'])
        r_qtxt.font.name = 'Calibri'
        r_qtxt.font.size = Pt(11)
        r_qtxt.font.bold = True
        r_qtxt.font.color.rgb = C_DARK

        for opt_letter, opt_text in item['opts']:
            p_opt = doc.add_paragraph()
            p_opt.paragraph_format.left_indent = Inches(0.28)
            p_opt.paragraph_format.space_before = Pt(0)
            p_opt.paragraph_format.space_after = Pt(2.5)
            p_opt.paragraph_format.line_spacing = 1.15

            r_opt_lbl = p_opt.add_run(f"{opt_letter})  ")
            r_opt_lbl.font.name = 'Calibri'
            r_opt_lbl.font.size = Pt(10.5)
            r_opt_lbl.font.bold = True
            r_opt_lbl.font.color.rgb = C_SLATE

            r_opt_val = p_opt.add_run(opt_text)
            r_opt_val.font.name = 'Calibri'
            r_opt_val.font.size = Pt(10.5)
            r_opt_val.font.color.rgb = C_BODY

    # Answer Key Table for Part 1
    p_ak1 = doc.add_paragraph()
    p_ak1.paragraph_format.space_before = Pt(18)
    p_ak1.paragraph_format.space_after = Pt(8)
    p_ak1.paragraph_format.keep_with_next = True
    r_ak1 = p_ak1.add_run("Answer Key: Part 1 (History Rebellions)")
    r_ak1.font.name = 'Calibri'
    r_ak1.font.size = Pt(13)
    r_ak1.font.bold = True
    r_ak1.font.color.rgb = C_NAVY

    # Create clean 4-column table for 20 answers (10 rows)
    # Col 0: Q#, Col 1: Answer, Col 2: Q#, Col 3: Answer
    table1 = doc.add_table(rows=11, cols=4)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table1, color="CBD5E1")

    # Header Row
    headers = ["Q#", "Answer", "Q#", "Answer"]
    hdr_cells = table1.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], HEX_NAVY)
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in [0, 2] else WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.name = 'Calibri'
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    # Set column widths: 0.6", 2.65", 0.6", 2.65" (Total 6.5" printable area)
    col_widths = [Inches(0.6), Inches(2.65), Inches(0.6), Inches(2.65)]
    for row in table1.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    # Populate 10 data rows
    for r_idx in range(10):
        row_cells = table1.rows[r_idx + 1].cells
        bg_color = HEX_ROW_EVEN if r_idx % 2 == 0 else HEX_ROW_ODD

        # Left pair: Q1 to Q10
        q_num_left = r_idx + 1
        ans_left = next((a[1] for a in ans1_list if a[0] == q_num_left), "")
        row_cells[0].text = str(q_num_left)
        row_cells[1].text = ans_left

        # Right pair: Q11 to Q20
        q_num_right = r_idx + 11
        ans_right = next((a[1] for a in ans1_list if a[0] == q_num_right), "")
        row_cells[2].text = str(q_num_right)
        row_cells[3].text = ans_right

        for c_idx in range(4):
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=70, bottom=70, left=100, right=100)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 2] else WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Calibri'
                r.font.size = Pt(9.5)
                if c_idx in [0, 2]:
                    r.font.bold = True
                    r.font.color.rgb = C_NAVY
                else:
                    r.font.color.rgb = C_DARK

    # Page Break before Part 2
    doc.add_page_break()

    # ==========================================
    # PART 2: 20 Questions on Indian Freedom Fighters
    # ==========================================
    p_p2 = doc.add_paragraph()
    p_p2.paragraph_format.space_before = Pt(10)
    p_p2.paragraph_format.space_after = Pt(4)
    p_p2.paragraph_format.keep_with_next = True
    r_p2 = p_p2.add_run("Part 2: 20 Questions on Indian Freedom Fighters")
    r_p2.font.name = 'Calibri'
    r_p2.font.size = Pt(15)
    r_p2.font.bold = True
    r_p2.font.color.rgb = C_NAVY

    p_desc2 = doc.add_paragraph()
    p_desc2.paragraph_format.space_before = Pt(0)
    p_desc2.paragraph_format.space_after = Pt(14)
    p_desc2.paragraph_format.keep_with_next = True
    r_desc2 = p_desc2.add_run(
        "Explore key personalities, pivotal movements, and significant milestones of the Indian Independence Struggle. "
        "Answers are detailed after each question and summarized in the answer key below."
    )
    r_desc2.font.name = 'Calibri'
    r_desc2.font.size = Pt(10)
    r_desc2.font.italic = True
    r_desc2.font.color.rgb = C_SLATE

    for item in q2_list:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.left_indent = Inches(0)
        p_q.paragraph_format.space_before = Pt(9)
        p_q.paragraph_format.space_after = Pt(3)
        p_q.paragraph_format.line_spacing = 1.15
        p_q.paragraph_format.keep_with_next = True

        r_qnum = p_q.add_run(f"{item['num']}. ")
        r_qnum.font.name = 'Calibri'
        r_qnum.font.size = Pt(11)
        r_qnum.font.bold = True
        r_qnum.font.color.rgb = C_NAVY

        r_qtxt = p_q.add_run(item['q'])
        r_qtxt.font.name = 'Calibri'
        r_qtxt.font.size = Pt(11)
        r_qtxt.font.bold = True
        r_qtxt.font.color.rgb = C_DARK

        for opt_letter, opt_text in item['opts']:
            p_opt = doc.add_paragraph()
            p_opt.paragraph_format.left_indent = Inches(0.28)
            p_opt.paragraph_format.space_before = Pt(0)
            p_opt.paragraph_format.space_after = Pt(2.5)
            p_opt.paragraph_format.line_spacing = 1.15

            r_opt_lbl = p_opt.add_run(f"{opt_letter})  ")
            r_opt_lbl.font.name = 'Calibri'
            r_opt_lbl.font.size = Pt(10.5)
            r_opt_lbl.font.bold = True
            r_opt_lbl.font.color.rgb = C_SLATE

            r_opt_val = p_opt.add_run(opt_text)
            r_opt_val.font.name = 'Calibri'
            r_opt_val.font.size = Pt(10.5)
            r_opt_val.font.color.rgb = C_BODY

        if item.get('ans'):
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.left_indent = Inches(0.28)
            p_ans.paragraph_format.space_before = Pt(2)
            p_ans.paragraph_format.space_after = Pt(8)
            p_ans.paragraph_format.line_spacing = 1.15

            r_ans_lbl = p_ans.add_run("Answer: ")
            r_ans_lbl.font.name = 'Calibri'
            r_ans_lbl.font.size = Pt(10)
            r_ans_lbl.font.bold = True
            r_ans_lbl.font.color.rgb = C_GREEN

            r_ans_val = p_ans.add_run(item['ans'])
            r_ans_val.font.name = 'Calibri'
            r_ans_val.font.size = Pt(10)
            r_ans_val.font.color.rgb = C_GREEN

    # Answer Key Table for Part 2
    p_ak2 = doc.add_paragraph()
    p_ak2.paragraph_format.space_before = Pt(18)
    p_ak2.paragraph_format.space_after = Pt(8)
    p_ak2.paragraph_format.keep_with_next = True
    r_ak2 = p_ak2.add_run("Answer Key: Part 2 (Indian Freedom Fighters)")
    r_ak2.font.name = 'Calibri'
    r_ak2.font.size = Pt(13)
    r_ak2.font.bold = True
    r_ak2.font.color.rgb = C_NAVY

    table2 = doc.add_table(rows=11, cols=4)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table2, color="CBD5E1")

    hdr_cells2 = table2.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells2[i].text = h
        set_cell_background(hdr_cells2[i], HEX_NAVY)
        set_cell_margins(hdr_cells2[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells2[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in [0, 2] else WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.name = 'Calibri'
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    for row in table2.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    for r_idx in range(10):
        row_cells = table2.rows[r_idx + 1].cells
        bg_color = HEX_ROW_EVEN if r_idx % 2 == 0 else HEX_ROW_ODD

        q_num_left = r_idx + 1
        q_item_left = next((item for item in q2_list if item['num'] == q_num_left), None)
        ans_left = q_item_left['ans'] if q_item_left and q_item_left.get('ans') else ""
        row_cells[0].text = str(q_num_left)
        row_cells[1].text = ans_left

        q_num_right = r_idx + 11
        q_item_right = next((item for item in q2_list if item['num'] == q_num_right), None)
        ans_right = q_item_right['ans'] if q_item_right and q_item_right.get('ans') else ""
        row_cells[2].text = str(q_num_right)
        row_cells[3].text = ans_right

        for c_idx in range(4):
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=70, bottom=70, left=100, right=100)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 2] else WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Calibri'
                r.font.size = Pt(9.5)
                if c_idx in [0, 2]:
                    r.font.bold = True
                    r.font.color.rgb = C_NAVY
                else:
                    r.font.color.rgb = C_DARK

    doc.save(output_path)
    print(f"Document successfully written to {output_path}")

if __name__ == '__main__':
    q1, a1, q2 = parse_original_document('backup/testing_ems_board/history_backup.docx')
    create_aligned_document(q1, a1, q2, 'backup/testing_ems_board/history.docx')
